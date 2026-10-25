"""CSV/XLSX timesheet reader: upload limits, formula-safe text, header signature.

This module only turns bytes into a raw table. It never interprets hours, dates
or money; later payroll steps map columns and build typed entries from it.
"""

from __future__ import annotations

import csv
import hashlib
import io
import zipfile
from collections.abc import Generator, Iterable, Sequence
from contextlib import closing
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import (  # type: ignore[import-untyped]
    InvalidFileException,
)

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500
MAX_RAW_ROWS = 5000
MAX_UNPACKED_BYTES = 20 * 1024 * 1024

_XLSX_MAGIC = b"PK\x03\x04"
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_DELIMITERS = ",;\t"


class TimesheetError(Exception):
    """A timesheet upload was rejected. `message` is plain English for the UI."""

    def __init__(
        self, code: str, message: str, details: dict[str, Any] | None = None
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.details = details


@dataclass(frozen=True)
class RawTable:
    """Headers plus non-blank data rows, with cells exactly as read.

    ``row_numbers[i]`` is the physical row of ``rows[i]`` in the upload (the header
    row is 1 and blank rows count). It defaults to 2, 3, ... for hand-built tables.
    """

    headers: tuple[str, ...]
    rows: tuple[tuple[object, ...], ...]
    row_numbers: tuple[int, ...] = ()

    def __post_init__(self) -> None:
        if not self.row_numbers:
            object.__setattr__(self, "row_numbers", tuple(range(2, len(self.rows) + 2)))
        elif len(self.row_numbers) != len(self.rows):
            msg = "row_numbers must have one entry per row"
            raise ValueError(msg)


def escape_cell(text: str) -> str:
    """Prefix `'` to text a spreadsheet could execute as a formula."""
    if text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text


def normalise_header(header: str) -> str:
    """Strip, collapse inner whitespace and casefold: the identity of a column."""
    return " ".join(header.split()).casefold()


def header_signature(headers: Sequence[str]) -> str:
    """sha256 over normalised headers (strip, collapse spaces, casefold)."""
    norm = [normalise_header(h) for h in headers]
    return hashlib.sha256("\x1f".join(norm).encode("utf-8")).hexdigest()


def file_hash(data: bytes) -> str:
    """sha256 of the raw upload bytes."""
    return hashlib.sha256(data).hexdigest()


def _is_blank(row: Iterable[object]) -> bool:
    return all(c is None or (isinstance(c, str) and not c.strip()) for c in row)


def _unsupported() -> TimesheetError:
    return TimesheetError("UNSUPPORTED_FILE_TYPE", "Upload a .csv or .xlsx timesheet.")


def _invalid(message: str) -> TimesheetError:
    return TimesheetError("TIMESHEET_INVALID", message)


def _read_csv_rows(data: bytes) -> Generator[list[object], None, None]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    sample = text[:4096]
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=_DELIMITERS).delimiter
    except csv.Error:
        # Sniffer gives up on small or uniform files; use the header line.
        first = next((ln for ln in sample.splitlines() if ln.strip()), "")
        best = max(_DELIMITERS, key=first.count)
        if first.count(best) > first.count(","):
            delimiter = best
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    for row in reader:
        yield [c if c != "" else None for c in row]


def _read_xlsx_rows(data: bytes) -> Generator[list[object], None, None]:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if sum(i.file_size for i in zf.infolist()) > MAX_UNPACKED_BYTES:
            raise TimesheetError(
                "FILE_TOO_LARGE", "This spreadsheet is too large once unpacked."
            )
    wb = load_workbook(
        io.BytesIO(data), read_only=True, data_only=False, keep_links=False
    )
    try:
        if not wb.worksheets:
            return
        ws = wb.worksheets[0]
        # The declared <dimension> can be wrong and would silently drop rows.
        ws.reset_dimensions()
        for row in ws.iter_rows(values_only=True):
            yield list(row)
    finally:
        wb.close()


def _parse(raw: Generator[list[object], None, None]) -> RawTable:
    """Consume `raw` lazily, stopping as soon as the row cap is exceeded."""
    headers: tuple[str, ...] | None = None
    width = 0
    rows: list[tuple[object, ...]] = []
    numbers: list[int] = []
    for number, row in enumerate(raw, start=1):
        if number > MAX_RAW_ROWS:
            raise _invalid("This timesheet has too many rows, including empty ones.")
        if _is_blank(row):
            continue
        if headers is None:
            cells = list(row)
            while cells and (cells[-1] is None or not str(cells[-1]).strip()):
                cells.pop()
            if any(c is None or not str(c).strip() for c in cells):
                raise _invalid(
                    "This timesheet needs a header row with a name for every column."
                )
            headers = tuple(str(c).strip() for c in cells)
            if len({normalise_header(h) for h in headers}) != len(headers):
                raise _invalid("Two columns have the same header name.")
            width = len(headers)
            continue
        if len(rows) >= MAX_ROWS:
            raise _invalid("This timesheet has more than 500 rows.")
        extra = row[width:]
        if not _is_blank(extra):
            raise TimesheetError(
                "TIMESHEET_INVALID",
                "A row has more cells than there are column headers.",
                {"rows": [{"row": number, "error": "more cells than headers"}]},
            )
        cells = list(row[:width])
        cells.extend([None] * (width - len(cells)))
        rows.append(tuple(cells))
        numbers.append(number)
    if headers is None:
        raise _invalid("This timesheet is empty.")
    return RawTable(headers=headers, rows=tuple(rows), row_numbers=tuple(numbers))


def read_table(data: bytes, filename: str) -> RawTable:
    """Read a CSV/XLSX upload. Checks run in order: size, type, parse."""
    if len(data) > MAX_BYTES:
        raise TimesheetError("FILE_TOO_LARGE", "This timesheet is larger than 2 MB.")

    name = filename.lower()
    if name.endswith(".xlsx"):
        if not data.startswith(_XLSX_MAGIC):
            raise _unsupported()
        raw = _read_xlsx_rows(data)
    elif name.endswith(".csv"):
        raw = _read_csv_rows(data)
    else:
        raise _unsupported()

    try:
        with closing(raw):
            return _parse(raw)
    except csv.Error as exc:
        raise _invalid("This timesheet could not be read as CSV.") from exc
    except (zipfile.BadZipFile, KeyError, InvalidFileException, OSError) as exc:
        raise _unsupported() from exc
