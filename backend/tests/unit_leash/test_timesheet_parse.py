"""Tests for the CSV/XLSX timesheet reader (plan WP-5.3, TST-17)."""

from __future__ import annotations

import hashlib
from collections.abc import Generator
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from backend.app.leash.payroll.parse import (
    MAX_BYTES,
    MAX_ROWS,
    TimesheetError,
    escape_cell,
    file_hash,
    header_signature,
    read_table,
)

DEMO_CSV = (
    b"Employee,Day,Hrs,Project\n"
    b"Omar Hassan,2026-10-01,8,MS-1\n"
    b"Omar Hassan,2026-10-02,7.5,MS-1\n"
    b"Sara Ali,2026-10-01,6,MS-1\n"
    b"Sara Ali,2026-10-02,6,MS-1\n"
    b"Youssef Adel,2026-10-01,4,MS-2\n"
    b"Youssef Adel,2026-10-02,5,MS-2\n"
    b"Omar Hassan,2026-10-03,3,MS-2\n"
    b"Sara Ali,2026-10-03,2.25,MS-2\n"
    b"Youssef Adel,2026-10-03,8,MS-2\n"
)


def _xlsx(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _code(excinfo: pytest.ExceptionInfo[TimesheetError]) -> str:
    return excinfo.value.code


def test_reads_demo_csv() -> None:
    table = read_table(DEMO_CSV, "timesheet.csv")
    assert table.headers == ("Employee", "Day", "Hrs", "Project")
    assert len(table.rows) == 9
    assert table.rows[0] == ("Omar Hassan", "2026-10-01", "8", "MS-1")


def test_reads_xlsx_with_dates_and_numbers() -> None:
    data = _xlsx(
        [
            ["Employee", "Day", "Hrs", "Project"],
            ["Omar Hassan", date(2026, 10, 1), 7.5, "MS-1"],
            ["Sara Ali", datetime(2026, 10, 2, 0, 0), 6, "MS-1"],
        ]
    )
    table = read_table(data, "ts.XLSX")
    assert table.headers == ("Employee", "Day", "Hrs", "Project")
    assert len(table.rows) == 2
    assert table.rows[0][1] in (date(2026, 10, 1), datetime(2026, 10, 1))
    assert table.rows[0][2] == 7.5
    assert table.rows[1][1] == datetime(2026, 10, 2)
    assert table.rows[1][2] == 6


def test_csv_dialects() -> None:
    text = (
        "\ufeffEmployee;Day;Hrs\r\nOmar;2026-10-01;8\r\nSara;2026-10-02;6\r\n\r\n\r\n"
    )
    table = read_table(text.encode("utf-8"), "a.csv")
    assert table.headers == ("Employee", "Day", "Hrs")
    assert table.rows == (("Omar", "2026-10-01", "8"), ("Sara", "2026-10-02", "6"))

    tabbed = read_table(b"A\tB\n1\t2\n", "a.csv")
    assert tabbed.headers == ("A", "B")
    assert tabbed.rows == (("1", "2"),)


def test_csv_cp1252_fallback() -> None:
    data = "Employee,Hrs\nJos\xe9,8\n".encode("cp1252")
    table = read_table(data, "a.csv")
    assert table.rows == (("Jos\xe9", "8"),)


def test_rejects_over_2_mb() -> None:
    assert MAX_BYTES == 2 * 1024 * 1024
    # Garbage bytes: would fail parsing, so FILE_TOO_LARGE proves size goes first.
    with pytest.raises(TimesheetError) as exc:
        read_table(b"\xff" * (MAX_BYTES + 1), "big.pdf")
    assert _code(exc) == "FILE_TOO_LARGE"
    assert str(exc.value) == "This timesheet is larger than 2 MB."


def test_rejects_over_500_rows() -> None:
    assert MAX_ROWS == 500
    header = "Employee,Day,Hrs\n"
    ok = header + "".join("Omar,2026-10-01,1\n" for _ in range(500))
    assert len(read_table(ok.encode(), "a.csv").rows) == 500
    bad = header + "".join("Omar,2026-10-01,1\n" for _ in range(501))
    with pytest.raises(TimesheetError) as exc:
        read_table(bad.encode(), "a.csv")
    assert _code(exc) == "TIMESHEET_INVALID"
    assert str(exc.value) == "This timesheet has more than 500 rows."


def test_blank_rows_do_not_count() -> None:
    body = "Omar,2026-10-01,1\n" * 500 + ",,\n   , ,\n" * 20
    table = read_table(("Employee,Day,Hrs\n" + body).encode(), "a.csv")
    assert len(table.rows) == 500


@pytest.mark.parametrize(
    ("data", "name"),
    [
        (b"%PDF-1.4", "a.pdf"),
        (b"PK\x03\x04", "a.xls"),
        (b"not a zip at all", "a.xlsx"),
    ],
)
def test_rejects_unsupported_type(data: bytes, name: str) -> None:
    with pytest.raises(TimesheetError) as exc:
        read_table(data, name)
    assert _code(exc) == "UNSUPPORTED_FILE_TYPE"


def test_rejects_corrupt_xlsx_zip() -> None:
    with pytest.raises(TimesheetError) as exc:
        read_table(b"PK\x03\x04garbage-not-really-a-zip", "a.xlsx")
    assert _code(exc) == "UNSUPPORTED_FILE_TYPE"


def test_rejects_empty_and_headerless() -> None:
    for data, name in [
        (b"", "a.csv"),
        (b"\n\n", "a.csv"),
        (b",Day,Hrs\nOmar,2026-10-01,1\n", "a.csv"),
        (_xlsx([]), "a.xlsx"),
        (_xlsx([[None, "Day"], ["x", "y"]]), "a.xlsx"),
    ]:
        with pytest.raises(TimesheetError) as exc:
            read_table(data, name)
        assert _code(exc) == "TIMESHEET_INVALID", name


@pytest.mark.parametrize(
    "text",
    ["=cmd|' /C calc'!A0", "+1", "-2+3", "@SUM(A1)", "\tx", "\rx"],
)
def test_formula_cells_are_escaped(text: str) -> None:
    assert escape_cell(text) == "'" + text


@pytest.mark.parametrize("text", ["Omar", "", "8.5", "a=b", " =x"])
def test_normal_text_is_unchanged(text: str) -> None:
    assert escape_cell(text) == text


def test_xlsx_formula_is_read_as_text() -> None:
    data = _xlsx([["Employee", "Hrs"], ["Omar", "=1+1"]])
    table = read_table(data, "a.xlsx")
    assert table.rows[0][1] == "=1+1"


def test_xlsx_headers_are_stripped_and_blank_rows_dropped() -> None:
    data = _xlsx([[" Employee ", 2], ["Omar", 1], [None, None], ["  ", None]])
    table = read_table(data, "a.xlsx")
    assert table.headers == ("Employee", "2")
    assert table.rows == (("Omar", 1),)


def test_header_signature() -> None:
    sig = header_signature(["Employee", "Day"])
    assert sig != header_signature(["Day", "Employee"])
    assert header_signature([" Employee ", "DAY"]) == header_signature(
        ["employee", "day"]
    )
    assert header_signature(["Hours  Worked"]) == header_signature(["hours worked"])
    assert len(sig) == 64
    assert sig == sig.lower()
    assert all(c in "0123456789abcdef" for c in sig)
    expected = hashlib.sha256(b"employee\x1fday").hexdigest()
    assert sig == expected


def test_file_hash() -> None:
    assert file_hash(DEMO_CSV) == hashlib.sha256(DEMO_CSV).hexdigest()


def _counting(monkeypatch: pytest.MonkeyPatch, name: str) -> list[int]:
    from backend.app.leash.payroll import parse

    original = getattr(parse, name)
    seen = [0]

    def wrapper(data: bytes) -> Generator[list[object], None, None]:
        for row in original(data):
            seen[0] += 1
            yield row

    monkeypatch.setattr(parse, name, wrapper)
    return seen


def test_xlsx_row_cap_stops_reading_early(monkeypatch: pytest.MonkeyPatch) -> None:
    rows: list[list[object]] = [["Employee", "Hrs"]]
    rows += [["Omar", 1] for _ in range(5000)]
    seen = _counting(monkeypatch, "_read_xlsx_rows")
    with pytest.raises(TimesheetError) as exc:
        read_table(_xlsx(rows), "a.xlsx")
    assert _code(exc) == "TIMESHEET_INVALID"
    assert seen[0] <= 502


def test_csv_row_cap_stops_reading_early(monkeypatch: pytest.MonkeyPatch) -> None:
    text = "Employee,Hrs\n" + "Omar,1\n" * 5000
    seen = _counting(monkeypatch, "_read_csv_rows")
    with pytest.raises(TimesheetError) as exc:
        read_table(text.encode(), "a.csv")
    assert _code(exc) == "TIMESHEET_INVALID"
    assert seen[0] <= 502


def test_oversized_csv_field_is_invalid_not_500() -> None:
    data = b"Employee,Hrs\n" + b"x" * 140_000 + b",1\n"
    assert len(data) < MAX_BYTES
    with pytest.raises(TimesheetError) as exc:
        read_table(data, "a.csv")
    assert _code(exc) == "TIMESHEET_INVALID"


def test_trailing_blank_headers_are_trimmed() -> None:
    csv_table = read_table(b"a,b,\n1,2,\n3,4,\n", "a.csv")
    assert csv_table.headers == ("a", "b")
    assert csv_table.rows == ((("1", "2")), ("3", "4"))
    xl = read_table(_xlsx([["a", "b", None, "  "], ["1", 2, None, None]]), "a.xlsx")
    assert xl.headers == ("a", "b")
    assert xl.rows == (("1", 2),)


def test_blank_header_in_the_middle_is_invalid() -> None:
    with pytest.raises(TimesheetError) as exc:
        read_table(b"a,,c\n1,2,3\n", "a.csv")
    assert _code(exc) == "TIMESHEET_INVALID"


def test_overflow_cells_nonblank_rejected_blank_dropped() -> None:
    ok = read_table(b"a,b\n1,2,,\n3,4,  \n", "a.csv")
    assert ok.rows == (("1", "2"), ("3", "4"))
    with pytest.raises(TimesheetError) as exc:
        read_table(b"a,b\n1,2\n3,4,oops\n", "a.csv")
    assert _code(exc) == "TIMESHEET_INVALID"
    assert exc.value.details is not None
    assert [r["row"] for r in exc.value.details["rows"]] == [3]


def _xlsx_with_blank_rows(blank: int) -> bytes:
    import zipfile

    src = zipfile.ZipFile(BytesIO(_xlsx([["Employee", "Hrs"]])))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            body = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                rows = "".join(f'<row r="{n}"/>' for n in range(2, blank + 2))
                body = body.replace(b'<dimension ref="A1:B1" />', b"")
                body = body.replace(b"</sheetData>", rows.encode() + b"</sheetData>")
            dst.writestr(info, body)
    return out.getvalue()


def test_xlsx_blank_rows_are_capped(monkeypatch: pytest.MonkeyPatch) -> None:
    from backend.app.leash.payroll.parse import MAX_RAW_ROWS

    assert MAX_RAW_ROWS == 5000
    data = _xlsx_with_blank_rows(20_000)
    seen = _counting(monkeypatch, "_read_xlsx_rows")
    with pytest.raises(TimesheetError) as exc:
        read_table(data, "a.xlsx")
    assert _code(exc) == "TIMESHEET_INVALID"
    assert str(exc.value) == "This timesheet has too many rows, including empty ones."
    assert seen[0] <= MAX_RAW_ROWS + 1


def test_csv_blank_lines_are_capped(monkeypatch: pytest.MonkeyPatch) -> None:
    from backend.app.leash.payroll.parse import MAX_RAW_ROWS

    data = b"Employee,Hrs\n" + b",\n" * 6000
    seen = _counting(monkeypatch, "_read_csv_rows")
    with pytest.raises(TimesheetError) as exc:
        read_table(data, "a.csv")
    assert _code(exc) == "TIMESHEET_INVALID"
    assert seen[0] <= MAX_RAW_ROWS + 1


def test_xlsx_unpacked_size_is_capped(monkeypatch: pytest.MonkeyPatch) -> None:
    import zipfile

    from backend.app.leash.payroll import parse

    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml", b"<row/>" * (21 * 1024 * 1024 // 6))
    data = buf.getvalue()
    assert len(data) < MAX_BYTES

    def boom(*args: object, **kwargs: object) -> None:
        raise AssertionError("load_workbook must not be called")

    monkeypatch.setattr(parse, "load_workbook", boom)
    with pytest.raises(TimesheetError) as exc:
        read_table(data, "a.xlsx")
    assert _code(exc) == "FILE_TOO_LARGE"
    assert str(exc.value) == "This spreadsheet is too large once unpacked."
